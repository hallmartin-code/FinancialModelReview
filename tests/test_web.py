"""Web app surface: health, gating, upload validation, job lifecycle."""

from __future__ import annotations

import time

import pytest
from fastapi.testclient import TestClient
from reportlab.pdfgen import canvas

from deckscan.web import app

client = TestClient(app)


@pytest.fixture
def deck_bytes(tmp_path):
    path = tmp_path / "deck.pdf"
    pdf = canvas.Canvas(str(path))
    pdf.drawString(72, 720, "Northwind Analytics — Seed round")
    pdf.save()
    return path.read_bytes()


def test_healthz_needs_no_key():
    response = client.get("/healthz")
    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] == "ok"
    assert payload["api_key_configured"] is False


def test_index_renders_and_warns_about_the_missing_key():
    response = client.get("/")
    assert response.status_code == 200
    assert "ANTHROPIC_API_KEY is not configured" in response.text
    assert "disabled" in response.text  # the CTA is inert without a key


def test_index_carries_the_brand_and_offers_every_supported_type():
    response = client.get("/")
    assert "Ten Capital" in response.text
    assert "brand-mark" in response.text
    for suffix in (".pdf", ".pptx", ".docx"):
        assert suffix in response.text
    for suffix in (".xlsx", ".xlsm", ".csv"):
        assert suffix in response.text


def test_index_claims_nothing_the_app_does_not_do():
    """The UI must not promise an email copy the backend never sends."""
    response = client.get("/")
    assert "@" not in response.text.split("<style>")[0] or "console.anthropic.com" in response.text
    assert "emailed" in response.text  # only as the explicit denial
    assert "nothing is emailed" in response.text.lower()


def test_upload_without_a_key_is_refused(deck_bytes):
    response = client.post("/jobs", files={"deck": ("deck.pdf", deck_bytes, "application/pdf")})
    assert response.status_code == 503
    assert "ANTHROPIC_API_KEY" in response.json()["detail"]


@pytest.fixture
def model_bytes():
    from io import BytesIO

    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Period", "Revenue"])
    sheet.append(["FY2025", 1200000])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_a_model_with_no_deck_is_accepted(monkeypatch, model_bytes):
    """A spreadsheet on its own is a complete run, not a missing deck."""
    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-ant-test")
    response = client.post(
        "/jobs",
        files={"model": ("model.xlsx", model_bytes, XLSX_MEDIA)},
        follow_redirects=False,
    )
    assert response.status_code == 303


def test_the_job_page_for_a_model_only_run_is_titled_by_the_model(monkeypatch, model_bytes):
    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-ant-test")
    created = client.post(
        "/jobs",
        files={"model": ("northwind-model.xlsx", model_bytes, XLSX_MEDIA)},
        follow_redirects=False,
    )
    page = client.get(created.headers["location"]).text
    assert "northwind-model.xlsx" in page


def test_an_upload_with_neither_file_is_rejected(monkeypatch):
    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-ant-test")
    response = client.post("/jobs", files={"deck": ("", b"", "application/pdf")})
    assert response.status_code == 400
    assert "this run had neither" in response.json()["detail"]


def test_a_spreadsheet_dropped_on_the_main_zone_is_routed_to_the_model(monkeypatch, model_bytes):
    """The first dropzone takes everything; an .xlsx there is a model, not an error."""
    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-ant-test")
    created = client.post(
        "/jobs",
        files={"deck": ("northwind.xlsx", model_bytes, XLSX_MEDIA)},
        follow_redirects=False,
    )
    assert created.status_code == 303
    state = client.get(f"/api{created.headers['location']}").json()
    assert state["deck"] is None
    assert state["model"] == "northwind.xlsx"


def test_two_spreadsheets_and_no_deck_are_rejected(monkeypatch, model_bytes):
    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-ant-test")
    response = client.post(
        "/jobs",
        files={
            "deck": ("a.xlsx", model_bytes, XLSX_MEDIA),
            "model": ("b.xlsx", model_bytes, XLSX_MEDIA),
        },
    )
    assert response.status_code == 400
    assert "Two spreadsheets" in response.json()["detail"]


def test_the_main_dropzone_advertises_spreadsheets(monkeypatch):
    """The accept filter must not hide .xlsx from the file picker."""
    page = client.get("/").text
    accept = page.split('name="deck"')[1].split('accept="')[1].split('"')[0]
    assert ".xlsx" in accept and ".pdf" in accept


def test_unsupported_deck_type_is_rejected(monkeypatch, deck_bytes):
    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-ant-test")
    response = client.post(
        "/jobs", files={"deck": ("deck.key", deck_bytes, "application/octet-stream")}
    )
    assert response.status_code == 400
    assert "Unsupported deck type" in response.json()["detail"]


def test_word_decks_are_accepted(monkeypatch, tmp_path):
    from docx import Document

    path = tmp_path / "deck.docx"
    document = Document()
    document.add_paragraph("Northwind Analytics — Seed round")
    document.save(str(path))

    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-ant-test")
    response = client.post(
        "/jobs",
        files={
            "deck": (
                "deck.docx",
                path.read_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        follow_redirects=False,
    )
    assert response.status_code == 303  # accepted; the job runs in the background


def test_empty_upload_is_rejected(monkeypatch):
    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-ant-test")
    response = client.post("/jobs", files={"deck": ("deck.pdf", b"", "application/pdf")})
    assert response.status_code == 400


def test_the_form_offers_the_model_as_a_standalone_source():
    """The main dropzone must say it takes a model, and require nothing."""
    page = client.get("/").text
    assert "drop a deck or model here" in page
    assert 'name="deck"' in page and "required" not in page.split('name="deck"')[1][:200]


def test_an_unreadable_deck_finishes_with_a_warning_not_a_clean_result(monkeypatch, deck_bytes):
    """A bad key degrades to a gaps-only report — the UI must say so."""
    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-ant-invalid")
    created = client.post(
        "/jobs",
        files={"deck": ("deck.pdf", deck_bytes, "application/pdf")},
        follow_redirects=False,
    )
    job_url = created.headers["location"]

    for _ in range(40):
        state = client.get(f"/api{job_url}").json()
        if state["status"] in {"done", "error"}:
            break
        time.sleep(0.25)

    assert state["status"] == "done"  # the run still produces a report
    assert state["warning"], "an unreadable deck must not look like a clean result"
    assert state["screen"] and state["onepager"]


def test_unknown_job_is_404():
    response = client.get("/api/jobs/does-not-exist")
    assert response.status_code == 404


def test_password_gate_when_app_password_is_set(monkeypatch):
    monkeypatch.setenv("APP_PASSWORD", "letmein")
    unauthorized = client.get("/")
    assert unauthorized.status_code == 401
    authorized = client.get("/", auth=("analyst", "letmein"))
    assert authorized.status_code == 200


def test_open_by_default(monkeypatch):
    monkeypatch.delenv("APP_PASSWORD", raising=False)
    assert client.get("/").status_code == 200


def test_favicon_png_is_served():
    response = client.get("/static/favicon.png")
    assert response.status_code == 200
    assert response.headers["content-type"] == "image/png"
    assert response.content[:8] == b"\x89PNG\r\n\x1a\n"


def test_every_linked_icon_exists():
    """A <link> pointing at a missing file is a 404 on every page load."""
    import re

    page = client.get("/").text
    hrefs = re.findall(r'<link[^>]+rel="(?:icon|apple-touch-icon)"[^>]+href="([^"]+)"', page)
    assert hrefs, "the page links no icons"
    for href in hrefs:
        assert client.get(href).status_code == 200, href


def test_bare_favicon_ico_path_resolves():
    response = client.get("/favicon.ico")
    assert response.status_code == 200
    assert response.content[:8] == b"\x89PNG\r\n\x1a\n"
